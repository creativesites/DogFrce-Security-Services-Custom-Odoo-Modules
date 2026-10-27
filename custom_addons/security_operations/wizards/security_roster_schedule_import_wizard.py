# -*- coding: utf-8 -*-
import base64
import csv
import io
import re
import difflib
import logging
from datetime import datetime, date, timedelta

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError

_logger = logging.getLogger(__name__)


class SecurityRosterScheduleImportWizard(models.TransientModel):
    _name = "security.roster.schedule.import.wizard"
    _description = "Import Excel / CSV Roster Shifts & Override Live Schedule"

    site_id = fields.Many2one(
        "security.client.site",
        string="Target Site (Optional)",
        help="If set, all imported rows will apply to this site. If left empty, site names will be detected automatically from the sheet or row data.",
    )
    date_from = fields.Date(
        string="Cycle Start Date",
        default=lambda self: self._default_date_from(),
        required=True,
    )
    date_to = fields.Date(
        string="Cycle End Date",
        default=lambda self: self._default_date_to(),
        required=True,
    )
    import_mode = fields.Selection(
        [
            ("file", "Upload Spreadsheet (Excel .xlsx / CSV)"),
            ("text", "Paste Roster Text / Table"),
        ],
        string="Input Method",
        default="file",
        required=True,
    )
    file_data = fields.Binary(
        string="Excel / CSV Roster File",
    )
    file_name = fields.Char(string="File Name")
    raw_text = fields.Text(string="Pasted Roster Text")

    override_assigned = fields.Boolean(
        string="Override Current Assignments",
        default=True,
        help="If checked, imported guards will replace currently assigned guards. If unchecked, only open / unassigned shifts will be filled.",
    )
    create_missing_slots = fields.Boolean(
        string="Create Missing Slots",
        default=True,
        help="If a shift date/time in the spreadsheet does not exist yet, generate a new roster slot automatically.",
    )
    auto_enroll_guard_pool = fields.Boolean(
        string="Auto-Enroll Guards to Site Pool",
        default=True,
        help="If checked, any imported guard assigned to shifts on a site who is not yet in the site's regular guard pool will be automatically added to that site's pool.",
    )
    default_shift_template_id = fields.Many2one(
        "security.shift.template",
        string="Default Shift Template",
        help="Used if the spreadsheet row does not specify Day or Night shift.",
    )

    state = fields.Selection(
        [
            ("upload", "Upload & Configure"),
            ("preview", "Review & Validate"),
            ("done", "Import Complete"),
        ],
        default="upload",
        required=True,
    )

    line_ids = fields.One2many(
        "security.roster.schedule.import.wizard.line",
        "wizard_id",
        string="Parsed Schedule Lines",
    )

    # Summary KPIs
    total_lines = fields.Integer(string="Total Shifts Parsed", compute="_compute_kpis")
    fill_gaps_count = fields.Integer(string="Gaps to Fill", compute="_compute_kpis")
    override_count = fields.Integer(string="Overrides", compute="_compute_kpis")
    unchanged_count = fields.Integer(string="Already Matching", compute="_compute_kpis")
    clash_count = fields.Integer(string="Clashes & Warnings", compute="_compute_kpis")
    unmatched_guards = fields.Integer(string="Unmatched Guards", compute="_compute_kpis")
    unmatched_sites = fields.Integer(string="Unmatched Sites", compute="_compute_kpis")

    @api.model
    def _default_date_from(self):
        today = fields.Date.today()
        if today.day >= 21:
            return today.replace(day=21)
        prev_month = today.replace(day=1) - timedelta(days=1)
        return prev_month.replace(day=21)

    @api.model
    def _default_date_to(self):
        start = self._default_date_from()
        # Next month 20th
        next_month = (start.replace(day=28) + timedelta(days=5)).replace(day=20)
        return next_month

    @api.depends("line_ids", "line_ids.action_type", "line_ids.clash_type", "line_ids.employee_id", "line_ids.site_id")
    def _compute_kpis(self):
        for wiz in self:
            wiz.total_lines = len(wiz.line_ids)
            wiz.fill_gaps_count = len(wiz.line_ids.filtered(lambda l: l.action_type == "fill"))
            wiz.override_count = len(wiz.line_ids.filtered(lambda l: l.action_type == "override"))
            wiz.unchanged_count = len(wiz.line_ids.filtered(lambda l: l.action_type == "keep"))
            wiz.clash_count = len(wiz.line_ids.filtered(lambda l: l.clash_type not in ("none", False)))
            wiz.unmatched_guards = len(wiz.line_ids.filtered(lambda l: not l.employee_id and l.raw_guard_name))
            wiz.unmatched_sites = len(wiz.line_ids.filtered(lambda l: not l.site_id and l.raw_site_name))

    # ─────────────────────────────────────────────────────────────
    # PARSING ENGINE
    # ─────────────────────────────────────────────────────────────
    def action_parse_and_preview(self):
        """Parse uploaded Excel / CSV or text and populate line_ids for manager review."""
        self.ensure_one()

        parsed_records = []
        if self.import_mode == "text":
            if not self.raw_text or not self.raw_text.strip():
                raise UserError(_("Please paste roster text or table rows."))
            parsed_records = self._parse_text_table(self.raw_text)
        else:
            if not self.file_data:
                raise UserError(_("Please select an Excel (.xlsx / .xls) or CSV file to upload."))
            try:
                file_bytes = base64.b64decode(self.file_data)
            except Exception as e:
                raise UserError(_("Could not decode file: %s") % str(e))

            fname = (self.file_name or "").lower()
            if fname.endswith(".csv"):
                parsed_records = self._parse_csv(file_bytes)
            else:
                parsed_records = self._parse_excel(file_bytes)

        if not parsed_records:
            raise UserError(_("No shift or guard rows could be extracted from the uploaded document. Please check the format."))

        # Clear existing lines
        self.line_ids.unlink()

        # Cache reference data for fast matching
        all_guards = self.env["hr.employee"].search([
            ("security_guard", "=", True),
            ("active", "=", True),
        ])
        all_sites = self.env["security.client.site"].search([("active", "=", True)])
        all_templates = self.env["security.shift.template"].search([])

        # Pre-fetch existing slots in the date range
        existing_slots = self.env["security.roster.slot"].search([
            ("shift_date", ">=", self.date_from),
            ("shift_date", "<=", self.date_to),
            ("state", "!=", "cancelled"),
        ])
        # Map: (site_id, shift_date_str, template_id) -> slot
        slot_map = {}
        for s in existing_slots:
            k = (s.site_id.id, str(s.shift_date), s.shift_template_id.id if s.shift_template_id else 0)
            slot_map[k] = s

        # Pre-fetch approved leaves
        approved_leaves = self.env["security.leave.request"].search([
            ("state", "=", "approved"),
            ("date_to", ">=", self.date_from),
            ("date_from", "<=", self.date_to),
        ]) if "security.leave.request" in self.env else self.env["hr.employee"]

        new_lines = []
        for rec in parsed_records:
            raw_site = rec.get("site") or (self.site_id.name if self.site_id else "")
            raw_guard = rec.get("guard") or ""
            shift_d = rec.get("date")
            raw_shift = rec.get("shift") or ""

            if not shift_d:
                continue

            # 1. Match Site
            site = self.site_id or self._match_site(raw_site, all_sites)

            # 2. Match Guard
            guard, guard_conf = self._match_guard(raw_guard, all_guards)

            # 3. Match Shift Template
            template = self._match_shift_template(raw_shift, all_templates) or self.default_shift_template_id

            # 4. Check existing slot
            slot_key = (site.id if site else 0, str(shift_d), template.id if template else 0)
            existing_slot = slot_map.get(slot_key)
            if not existing_slot and site:
                # Try fallback matching site and date alone if template wasn't exact
                for s in existing_slots.filtered(lambda x: x.site_id.id == site.id and x.shift_date == shift_d):
                    existing_slot = s
                    break

            # 5. Determine action & conflicts
            action_type = "create"
            clash_type = "none"
            clash_detail = ""

            if not site:
                clash_type = "no_site"
                clash_detail = f"Site '{raw_site}' could not be matched"
            elif not guard:
                clash_type = "no_guard"
                clash_detail = f"Guard '{raw_guard}' could not be matched"
            elif guard.security_disqualified or not guard.active:
                clash_type = "disqualified"
                clash_detail = f"Guard {guard.name} is inactive or disqualified"
            else:
                # Check approved leave on shift_d
                on_leave = False
                if approved_leaves:
                    on_leave = bool(approved_leaves.filtered(
                        lambda l: l.employee_id.id == guard.id and l.date_from <= shift_d <= l.date_to
                    ))
                if on_leave:
                    clash_type = "leave"
                    clash_detail = f"{guard.name} has approved leave on {shift_d}"

            if existing_slot:
                if not existing_slot.employee_id:
                    action_type = "fill"
                elif existing_slot.employee_id.id == (guard.id if guard else 0):
                    action_type = "keep"
                else:
                    action_type = "override"
                    if clash_type == "none":
                        clash_detail = f"Replaces currently assigned guard: {existing_slot.employee_id.name}"
            else:
                action_type = "create"

            new_lines.append({
                "wizard_id": self.id,
                "shift_date": shift_d,
                "site_id": site.id if site else False,
                "raw_site_name": raw_site,
                "shift_template_id": template.id if template else False,
                "raw_shift": raw_shift,
                "employee_id": guard.id if guard else False,
                "raw_guard_name": raw_guard,
                "existing_slot_id": existing_slot.id if existing_slot else False,
                "existing_employee_id": existing_slot.employee_id.id if existing_slot and existing_slot.employee_id else False,
                "action_type": action_type,
                "clash_type": clash_type,
                "clash_detail": clash_detail,
                "apply_line": bool(site and guard and clash_type != "disqualified"),
            })

        # Check for intra-import double-booking across lines
        guard_dates = {}
        for l in new_lines:
            gid = l.get("employee_id")
            sd = l.get("shift_date")
            if gid and sd:
                guard_dates.setdefault((gid, str(sd)), []).append(l)

        for (gid, sd), entries in guard_dates.items():
            if len(entries) > 1:
                for entry in entries:
                    if entry["clash_type"] == "none":
                        entry["clash_type"] = "double_booking"
                        entry["clash_detail"] = f"Warning: Multiple shifts on {sd} in this import"

        self.env["security.roster.schedule.import.wizard.line"].create(new_lines)
        self.state = "preview"

        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    # ─────────────────────────────────────────────────────────────
    # COMMIT & OVERRIDE
    # ─────────────────────────────────────────────────────────────
    def action_apply_import(self):
        """Commit valid lines into the live roster slots."""
        self.ensure_one()

        valid_lines = self.line_ids.filtered(lambda l: l.apply_line and l.site_id and l.employee_id and l.shift_date)
        if not valid_lines:
            raise UserError(_("No lines are selected or valid to apply."))

        updated_slots = 0
        created_slots = 0
        filled_slots = 0
        overridden_slots = 0

        # Group lines by site and cycle to find or create batches
        sites = valid_lines.mapped("site_id")
        batch_model = self.env["security.roster.batch"]
        slot_model = self.env["security.roster.slot"]

        site_batches = {}
        for s in sites:
            b = batch_model.search([
                ("site_id", "=", s.id),
                ("date_from", "<=", self.date_from),
                ("date_to", ">=", self.date_to),
                ("state", "!=", "cancelled"),
            ], limit=1)
            if not b:
                # Find matching cycle batch
                b = batch_model.search([
                    ("site_id", "=", s.id),
                    ("date_from", "=", self.date_from),
                    ("date_to", "=", self.date_to),
                ], limit=1)
            if not b:
                # Create batch
                b = batch_model.create({
                    "name": f"{s.name}: {self.date_from} to {self.date_to}",
                    "site_id": s.id,
                    "partner_id": s.partner_id.id if s.partner_id else False,
                    "date_from": self.date_from,
                    "date_to": self.date_to,
                    "state": "generated",
                })
            site_batches[s.id] = b

        source_info = f"Excel Roster Import ({self.file_name or 'Manual File'})"

        # Track 1: Pool Auto-Learner: Auto-enroll imported guards to site regular guard pool
        enrolled_guards_by_site = {}
        if self.auto_enroll_guard_pool:
            for s in sites:
                site_guards = valid_lines.filtered(lambda l: l.site_id.id == s.id).mapped("employee_id")
                new_guards = site_guards - s.site_guard_pool_ids
                if new_guards:
                    s.write({
                        "site_guard_pool_ids": [(4, g.id) for g in new_guards],
                        "guard_pool_source": s.guard_pool_source or f"Auto-learned from {self.file_name or 'Excel Import'}",
                        "guard_pool_source_date": s.guard_pool_source_date or fields.Date.today(),
                    })
                    enrolled_guards_by_site[s.id] = new_guards
                    if hasattr(s, "message_post"):
                        s.message_post(
                            body=_(
                                "<b>Site Pool Auto-Learner:</b> Auto-enrolled %d guard(s) from roster import (%s): %s"
                            ) % (
                                len(new_guards),
                                self.file_name or "Spreadsheet",
                                ", ".join(new_guards.mapped("name")),
                            ),
                            message_type="notification",
                        )

        for line in valid_lines:
            target_slot = line.existing_slot_id
            batch = site_batches.get(line.site_id.id)

            if target_slot:
                was_empty = not bool(target_slot.employee_id)
                vals = {
                    "employee_id": line.employee_id.id,
                    "state": "assigned",
                    "critical_gap": False,
                }
                if not was_empty and target_slot.employee_id.id != line.employee_id.id:
                    vals["is_override"] = True
                    vals["override_reason"] = f"{source_info} - Replaced {target_slot.employee_id.name}"
                    overridden_slots += 1
                elif was_empty:
                    filled_slots += 1

                # Relief flag check
                eff_pool = line.site_id.site_guard_pool_ids
                if eff_pool and line.employee_id.id not in eff_pool.ids:
                    vals["is_relief"] = True
                    vals["relief_reason"] = f"Relief guard assigned via {source_info}"
                elif not eff_pool:
                    vals["is_relief"] = True
                    vals["relief_reason"] = f"Assigned via {source_info} (no regular site pool)"
                else:
                    vals["is_relief"] = False
                    vals["relief_reason"] = False

                target_slot.write(vals)
                updated_slots += 1
            elif self.create_missing_slots and batch:
                # Create new slot
                vals = {
                    "batch_id": batch.id,
                    "site_id": line.site_id.id,
                    "shift_date": line.shift_date,
                    "shift_template_id": line.shift_template_id.id if line.shift_template_id else False,
                    "employee_id": line.employee_id.id,
                    "state": "assigned",
                    "critical_gap": False,
                    "is_override": bool(line.clash_type != "none"),
                    "override_reason": f"{source_info} (Manually Created Slot)",
                }
                eff_pool = line.site_id.site_guard_pool_ids
                if eff_pool and line.employee_id.id not in eff_pool.ids:
                    vals["is_relief"] = True
                    vals["relief_reason"] = f"Relief guard assigned via {source_info}"
                elif not eff_pool:
                    vals["is_relief"] = True
                    vals["relief_reason"] = f"Assigned via {source_info} (no regular site pool)"
                else:
                    vals["is_relief"] = False
                    vals["relief_reason"] = False

                slot_model.create(vals)
                created_slots += 1
                filled_slots += 1

        self.state = "done"

        total_enrolled = sum(len(g_set) for g_set in enrolled_guards_by_site.values())
        msg = (
            f"Manual Excel Roster successfully imported!\n"
            f"• {updated_slots} existing shifts updated\n"
            f"• {created_slots} new shifts created\n"
            f"• {filled_slots} open gaps filled\n"
            f"• {overridden_slots} guard assignments overridden."
        )
        if total_enrolled:
            msg += f"\n• {total_enrolled} guard(s) auto-enrolled to site regular pool(s)."

        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("Roster Imported Successfully"),
                "message": msg,
                "type": "success",
                "sticky": True,
            },
        }

    def action_back_to_upload(self):
        self.ensure_one()
        self.line_ids.unlink()
        self.state = "upload"
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    # ─────────────────────────────────────────────────────────────
    # EXCEL / CSV / TEXT PARSERS
    # ─────────────────────────────────────────────────────────────
    def _parse_excel(self, file_bytes):
        """Parse Excel workbook supporting both tabular lists and date-column matrix sheets."""
        records = []
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except Exception as e:
            _logger.warning("openpyxl failed to read workbook: %s", e)
            return self._parse_csv(file_bytes)

        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            # Detect format: Matrix vs Tabular
            sheet_records = self._detect_and_parse_sheet(rows, default_site=sheetname)
            records.extend(sheet_records)

        return records

    def _detect_and_parse_sheet(self, rows, default_site=""):
        """Inspect sheet header row to determine if it's a date-column matrix or standard table."""
        if not rows:
            return []

        # Find first non-empty header row
        header_idx = 0
        header_row = None
        for i, r in enumerate(rows[:10]):
            non_empty = [c for c in r if c is not None and str(c).strip()]
            if len(non_empty) >= 3:
                header_idx = i
                header_row = [str(c or "").strip() for c in r]
                break

        if not header_row:
            return []

        # Check if columns contain dates (Matrix format: e.g. 21, 22, 23, or 2026-10-21)
        date_cols = {}
        for col_idx, col_name in enumerate(header_row):
            parsed_d = self._parse_date_header(col_name)
            if parsed_d:
                date_cols[col_idx] = parsed_d

        if len(date_cols) >= 3:
            # MATRIX FORMAT: Dates are across columns!
            return self._parse_matrix_format(rows[header_idx + 1:], header_row, date_cols, default_site)
        else:
            # TABULAR FORMAT: Standard column list
            return self._parse_tabular_format(rows[header_idx:], default_site)

    def _parse_matrix_format(self, data_rows, header_row, date_cols, default_site):
        """
        Matrix format:
        Row specifies: Site / Post / Shift (or Guard name).
        Columns 21, 22, 23... Contain Guard name (or Shift type if row is Guard).
        """
        records = []

        # Find metadata columns in left-hand headers
        site_col = None
        post_col = None
        shift_col = None
        guard_col = None

        for idx, col_name in enumerate(header_row):
            if idx in date_cols:
                continue
            cn = col_name.lower()
            if any(k in cn for k in ("site", "client", "location", "post", "station")):
                site_col = idx
            elif any(k in cn for k in ("shift", "template", "time", "day/night")):
                shift_col = idx
            elif any(k in cn for k in ("guard", "employee", "officer", "name", "personnel")):
                guard_col = idx

        for r in data_rows:
            if not any(c is not None for c in r):
                continue

            row_site = str(r[site_col]).strip() if site_col is not None and len(r) > site_col and r[site_col] else default_site
            row_shift = str(r[shift_col]).strip() if shift_col is not None and len(r) > shift_col and r[shift_col] else ""
            row_guard = str(r[guard_col]).strip() if guard_col is not None and len(r) > guard_col and r[guard_col] else ""

            # For each date column
            for col_idx, shift_date in date_cols.items():
                if col_idx >= len(r) or r[col_idx] is None:
                    continue
                cell_val = str(r[col_idx]).strip()
                if not cell_val or cell_val.lower() in ("off", "o", "rest", "r", "-", "0", "none"):
                    continue

                # If cell is guard name:
                guard_name = cell_val
                shift_type = row_shift
                # If row had guard name and cell has shift (e.g. "D", "N", "Day", "Night"):
                if row_guard and cell_val.lower() in ("d", "day", "n", "night", "ds", "ns"):
                    guard_name = row_guard
                    shift_type = cell_val

                records.append({
                    "site": row_site,
                    "date": shift_date,
                    "shift": shift_type,
                    "guard": guard_name,
                })

        return records

    def _parse_tabular_format(self, rows, default_site):
        """Tabular format: columns for Site, Date, Shift, Guard."""
        records = []
        headers = [str(c or "").strip().lower() for c in rows[0]]

        site_idx = self._find_col(headers, ("site", "client", "station", "location"))
        date_idx = self._find_col(headers, ("date", "shift date", "day"))
        shift_idx = self._find_col(headers, ("shift", "template", "time", "day/night"))
        guard_idx = self._find_col(headers, ("guard", "employee", "name", "officer", "id"))

        for r in rows[1:]:
            if not any(c is not None for c in r):
                continue
            raw_d = r[date_idx] if date_idx is not None and len(r) > date_idx else None
            parsed_d = self._parse_any_date(raw_d)
            if not parsed_d:
                continue

            site_val = str(r[site_idx]).strip() if site_idx is not None and len(r) > site_idx and r[site_idx] else default_site
            shift_val = str(r[shift_idx]).strip() if shift_idx is not None and len(r) > shift_idx and r[shift_idx] else ""
            guard_val = str(r[guard_idx]).strip() if guard_idx is not None and len(r) > guard_idx and r[guard_idx] else ""

            if guard_val:
                records.append({
                    "site": site_val,
                    "date": parsed_d,
                    "shift": shift_val,
                    "guard": guard_val,
                })

        return records

    def _parse_csv(self, file_bytes):
        """Parse CSV file in either tabular or matrix format."""
        text = ""
        for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                text = file_bytes.decode(enc)
                break
            except Exception:
                continue
        if not text:
            return []

        lines = [r for r in csv.reader(io.StringIO(text))]
        return self._detect_and_parse_sheet(lines, default_site=self.site_id.name if self.site_id else "")

    def _parse_text_table(self, raw_text):
        """Parse pasted text tables."""
        lines = []
        for line in raw_text.splitlines():
            line = line.strip()
            if not line:
                continue
            tokens = re.split(r"[,\t|]", line)
            lines.append([t.strip() for t in tokens])
        return self._detect_and_parse_sheet(lines, default_site=self.site_id.name if self.site_id else "")

    # ─────────────────────────────────────────────────────────────
    # MATCHING HELPERS
    # ─────────────────────────────────────────────────────────────
    def _find_col(self, headers, candidates):
        for idx, h in enumerate(headers):
            for c in candidates:
                if c in h:
                    return idx
        return None

    def _parse_date_header(self, col_text):
        """Detect if column header represents a date (e.g. 21, 22, 2026-10-21)."""
        if not col_text:
            return None
        text = str(col_text).strip()
        # Full date YYYY-MM-DD or DD/MM/YYYY
        parsed = self._parse_any_date(text)
        if parsed:
            return parsed

        # Number alone (e.g. 21, 22, 23)
        if text.isdigit():
            day_num = int(text)
            if 1 <= day_num <= 31:
                # Infer month and year from wizard cycle dates
                # If day_num >= 21, it belongs to date_from month; if <= 20, to date_to month
                base_d = self.date_from if day_num >= 21 else self.date_to
                try:
                    return base_d.replace(day=day_num)
                except ValueError:
                    return None
        return None

    def _parse_any_date(self, val):
        if not val:
            return None
        if isinstance(val, date):
            return val
        if isinstance(val, datetime):
            return val.date()
        s = str(val).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y", "%d %b", "%Y-%m-%d %H:%M:%S"):
            try:
                dt = datetime.strptime(s, fmt)
                if fmt in ("%d %b",):
                    # Add current year
                    dt = dt.replace(year=self.date_from.year)
                return dt.date()
            except ValueError:
                continue
        return None

    def _match_site(self, raw_site_name, all_sites):
        if not raw_site_name:
            return False
        clean = self._clean_name(raw_site_name)
        # 1. Exact
        for s in all_sites:
            if self._clean_name(s.name) == clean or self._clean_name(s.partner_id.name if s.partner_id else "") == clean:
                return s
        # 2. Substring
        for s in all_sites:
            s_clean = self._clean_name(s.name)
            if clean in s_clean or s_clean in clean:
                return s
        # 3. Fuzzy
        best_site = None
        best_ratio = 0.0
        for s in all_sites:
            ratio = difflib.SequenceMatcher(None, clean, self._clean_name(s.name)).ratio()
            if ratio > best_ratio:
                best_ratio = ratio
                best_site = s
        if best_site and best_ratio >= 0.75:
            return best_site
        return False

    def _match_guard(self, raw_guard_name, all_guards):
        if not raw_guard_name:
            return False, 0
        clean = self._clean_name(raw_guard_name)

        # 1. Code / ID match
        for g in all_guards:
            if g.employee_code and g.employee_code.strip().lower() == clean:
                return g, 100

        # 2. Exact name
        for g in all_guards:
            if self._clean_name(g.name) == clean:
                return g, 100

        # 3. Transposed name
        raw_parts = set(clean.split())
        for g in all_guards:
            g_parts = set(self._clean_name(g.name).split())
            if raw_parts and raw_parts == g_parts:
                return g, 95

        # 4. Fuzzy & Token-sorted fuzzy (handles reversed surname/first name + typos)
        clean_sorted = " ".join(sorted(clean.split()))
        best_guard = None
        best_ratio = 0.0
        for g in all_guards:
            g_clean = self._clean_name(g.name)
            r1 = difflib.SequenceMatcher(None, clean, g_clean).ratio()
            g_sorted = " ".join(sorted(g_clean.split()))
            r2 = difflib.SequenceMatcher(None, clean_sorted, g_sorted).ratio()
            ratio = max(r1, r2)
            if ratio > best_ratio:
                best_ratio = ratio
                best_guard = g
        if best_guard and best_ratio >= 0.78:
            return best_guard, int(best_ratio * 100)

        return False, 0

    def _match_shift_template(self, raw_shift, all_templates):
        if not raw_shift:
            return False
        clean = str(raw_shift).lower().strip()
        if clean in ("d", "day", "ds", "morning", "day shift"):
            for t in all_templates:
                if "day" in t.name.lower() or t.start_hour == 6.0:
                    return t
        elif clean in ("n", "night", "ns", "evening", "night shift"):
            for t in all_templates:
                if "night" in t.name.lower() or t.start_hour == 18.0:
                    return t
        for t in all_templates:
            if clean in t.name.lower():
                return t
        return False

    def _clean_name(self, text):
        if not text:
            return ""
        t = str(text).lower().strip()
        t = re.sub(r"^(guard|officer|mr|ms|mrs|site)\.?\s+", "", t)
        t = re.sub(r"[^\w\s]", " ", t)
        return " ".join(t.split())


class SecurityRosterScheduleImportWizardLine(models.TransientModel):
    _name = "security.roster.schedule.import.wizard.line"
    _description = "Parsed Roster Schedule Import Line"
    _order = "shift_date asc, raw_site_name asc, id asc"

    wizard_id = fields.Many2one("security.roster.schedule.import.wizard", required=True, ondelete="cascade")
    shift_date = fields.Date(string="Shift Date", required=True)
    site_id = fields.Many2one("security.client.site", string="Matched Site")
    raw_site_name = fields.Char(string="Site in File")
    shift_template_id = fields.Many2one("security.shift.template", string="Shift Template")
    raw_shift = fields.Char(string="Shift in File")
    employee_id = fields.Many2one("hr.employee", string="Matched Guard")
    raw_guard_name = fields.Char(string="Guard in File", required=True)

    existing_slot_id = fields.Many2one("security.roster.slot", string="Current Slot")
    existing_employee_id = fields.Many2one("hr.employee", string="Currently Assigned")

    action_type = fields.Selection(
        [
            ("fill", "Fill Gap"),
            ("override", "Override"),
            ("keep", "Keep Matching"),
            ("create", "Create Slot"),
        ],
        string="Action",
        required=True,
    )
    clash_type = fields.Selection(
        [
            ("none", "Clean"),
            ("double_booking", "Double-Booking Warning"),
            ("leave", "On Leave"),
            ("disqualified", "Disqualified/Inactive"),
            ("no_guard", "Unmatched Guard"),
            ("no_site", "Unmatched Site"),
        ],
        string="Conflict Status",
        default="none",
    )
    clash_detail = fields.Char(string="Warning / Notes")
    apply_line = fields.Boolean(string="Apply", default=True)
